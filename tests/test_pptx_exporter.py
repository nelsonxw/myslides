"""
Unit tests for PPTXExporter.
"""
import pytest
import tempfile
from pathlib import Path
from unittest.mock import Mock, MagicMock, patch
from pptx import Presentation

from myslides.generation.pptx_exporter import PPTXExporter


class TestPPTXExporter:
    """Test PPTXExporter class."""
    
    def test_initialization(self):
        """Test PPTXExporter initialization."""
        presentation = Presentation()
        exporter = PPTXExporter(presentation)
        assert exporter.presentation is presentation
    
    @patch('pathlib.Path.mkdir')
    def test_save(self, mock_mkdir):
        """Test saving presentation to file."""
        presentation = Presentation()
        # Mock the save method on the instance
        presentation.save = Mock()
        exporter = PPTXExporter(presentation)
        
        output_path = "test_output/generated.pptx"
        result = exporter.save(output_path)
        
        mock_mkdir.assert_called_once_with(parents=True, exist_ok=True)
        presentation.save.assert_called_once()
        assert result is not None
    
    @patch('pathlib.Path.mkdir')
    def test_save_creates_directory(self, mock_mkdir):
        """Test that save creates parent directories."""
        presentation = Presentation()
        presentation.save = Mock()
        exporter = PPTXExporter(presentation)
        
        output_path = "deep/nested/path/generated.pptx"
        exporter.save(output_path)
        
        mock_mkdir.assert_called_once_with(parents=True, exist_ok=True)
    
    @patch('tempfile.gettempdir')
    @patch('pathlib.Path.mkdir')
    def test_save_to_temp(self, mock_mkdir, mock_gettempdir):
        """Test saving presentation to temp directory."""
        mock_gettempdir.return_value = "/tmp"
        
        presentation = Presentation()
        presentation.save = Mock()
        exporter = PPTXExporter(presentation)
        
        result = exporter.save_to_temp("test.pptx")
        
        presentation.save.assert_called_once()
        assert "test.pptx" in result
    
    @patch('tempfile.gettempdir')
    @patch('pathlib.Path.mkdir')
    def test_save_to_temp_default_filename(self, mock_mkdir, mock_gettempdir):
        """Test saving to temp with default filename."""
        mock_gettempdir.return_value = "/tmp"
        
        presentation = Presentation()
        presentation.save = Mock()
        exporter = PPTXExporter(presentation)
        
        result = exporter.save_to_temp()
        
        presentation.save.assert_called_once()
        assert "generated_slide.pptx" in result
    
    def test_validate_valid_presentation(self):
        """Test validating a valid presentation."""
        presentation = Presentation()
        # Add a slide
        presentation.slides.add_slide(presentation.slide_layouts[0])
        
        exporter = PPTXExporter(presentation)
        result = exporter.validate()
        
        assert result is True
    
    def test_validate_empty_presentation(self):
        """Test validating an empty presentation."""
        presentation = Presentation()
        # No slides added
        
        exporter = PPTXExporter(presentation)
        result = exporter.validate()
        
        assert result is False
    
    def test_validate_with_exception(self):
        """Test validation when presentation raises exception."""
        presentation = Mock()
        presentation.slides = Mock(side_effect=Exception("Test error"))
        
        exporter = PPTXExporter(presentation)
        result = exporter.validate()
        
        assert result is False
    
    def test_create_empty(self):
        """Test creating an empty presentation."""
        presentation = PPTXExporter.create_empty()
        
        assert presentation is not None
        # Check it's a Presentation object by checking it has slides attribute
        assert hasattr(presentation, 'slides')
    
    @patch('myslides.generation.pptx_exporter.Presentation')
    def test_load_from_file_success(self, mock_presentation):
        """Test loading presentation from file successfully."""
        mock_presentation.return_value = Mock()
        
        result = PPTXExporter.load_from_file("test.pptx")
        
        mock_presentation.assert_called_once_with("test.pptx")
        assert result is not None
    
    @patch('pptx.Presentation')
    def test_load_from_file_failure(self, mock_presentation):
        """Test loading presentation from file when it fails."""
        mock_presentation.side_effect = Exception("File not found")
        
        result = PPTXExporter.load_from_file("nonexistent.pptx")
        
        assert result is None
