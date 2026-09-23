"""
Unit tests for SlideGenerator.
"""
import pytest
from pptx import Presentation

from myslides.generation.slide_generator import SlideGenerator
from myslides.generation.generation_models import (
    GenerationIntent, GenerationConfig, SlideType, 
    TextContent, ProcessStep, TimelineEvent
)


class TestSlideGenerator:
    """Test SlideGenerator class."""
    
    def test_initialization_default(self):
        """Test SlideGenerator initialization with default config."""
        generator = SlideGenerator()
        assert generator.config is not None
        assert generator.presentation is not None
        assert generator.style_inheritor is not None
        assert generator.chart_generator is not None
        assert generator.diagram_generator is not None
    
    def test_initialization_custom_config(self):
        """Test SlideGenerator initialization with custom config and presentation."""
        custom_p = Presentation()
        config = GenerationConfig(
            template_id="template_123",
            template_color_palette=["#1F4E79", "#2CA02C"],
            output_path="custom.pptx"
        )
        generator = SlideGenerator(config, presentation=custom_p)
        assert generator.config.template_id == "template_123"
        assert len(generator.config.template_color_palette) == 2
        assert generator.presentation is custom_p
    
    def test_generate_slide_title_slide(self):
        """Test generating a title slide."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.TITLE_SLIDE,
            content={
                "title": "My Presentation Title",
                "subtitle": "Subtitle Description"
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        # Title and subtitle textboxes
        assert len(presentation.slides[0].shapes) == 2
        texts = [s.text_frame.text for s in presentation.slides[0].shapes if s.has_text_frame]
        assert "My Presentation Title" in texts
        assert "Subtitle Description" in texts
    
    def test_generate_slide_process_flow(self):
        """Test generating a process flow slide end-to-end."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.PROCESS_FLOW,
            content={
                "title": "Our Process",
                "steps": ["Step 1", "Step 2", "Step 3"]
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        slide = presentation.slides[0]
        # Title textbox (1) + 3 step shapes + 2 arrows = 6 shapes
        assert len(slide.shapes) == 6
    
    def test_generate_slide_timeline(self):
        """Test generating a timeline slide end-to-end."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.TIMELINE,
            content={
                "title": "Project Timeline",
                "events": [
                    {"date": "2024-01-01", "title": "Start"},
                    {"date": "2024-06-01", "title": "Milestone"}
                ]
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        slide = presentation.slides[0]
        # Title (1) + line (1) + 2 markers (2) + 2 date boxes (2) + 2 title boxes (2) = 8 shapes
        assert len(slide.shapes) == 8
    
    def test_generate_slide_data_chart(self):
        """Test generating a data chart slide end-to-end."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.DATA_CHART,
            content={
                "title": "Sales Performance",
                "chart_data": {
                    "chart_type": "bar_vertical",
                    "title": "Q1-Q4 Revenue",
                    "categories": ["Q1", "Q2", "Q3", "Q4"],
                    "series_data": {"Revenue": [10.0, 15.0, 20.0, 25.0]}
                }
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        slide = presentation.slides[0]
        # Title textbox (1) + chart shape (1) = 2 shapes
        assert len(slide.shapes) == 2
    
    def test_generate_slide_text_heavy(self):
        """Test generating a text-heavy slide with bullets."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.TEXT_HEAVY,
            content={
                "title": "Key Points",
                "bullet_points": ["Point 1", "Point 2", "Point 3"]
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        slide = presentation.slides[0]
        assert len(slide.shapes) == 2
    
    def test_generate_slide_body_text(self):
        """Test generating slide with body text instead of bullets."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.TEXT_HEAVY,
            content={
                "title": "Overview",
                "body_text": "This is a detailed paragraph explaining the project scope."
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        slide = presentation.slides[0]
        assert len(slide.shapes) == 2
        body_box = slide.shapes[1]
        assert "This is a detailed paragraph" in body_box.text_frame.text
    
    def test_generate_slide_comparison(self):
        """Test generating a two-column comparison slide."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.COMPARISON,
            content={
                "title": "Comparison Matrix",
                "bullet_points": ["Plan A: Basic", "Plan A: Free", "Plan B: Advanced", "Plan B: Paid"]
            }
        )
        
        presentation = generator.generate_slide(intent)
        assert len(presentation.slides) == 1
        slide = presentation.slides[0]
        # Title (1) + Left column (1) + Right column (1) = 3 shapes
        assert len(slide.shapes) == 3
    
    def test_generate_deck(self):
        """Test generating a full multi-slide deck."""
        generator = SlideGenerator()
        intents = [
            GenerationIntent(slide_type=SlideType.TITLE_SLIDE, content={"title": "Deck Title"}),
            GenerationIntent(slide_type=SlideType.TEXT_HEAVY, content={"title": "Summary", "bullet_points": ["Point 1"]}),
            GenerationIntent(slide_type=SlideType.PROCESS_FLOW, content={"title": "Steps", "steps": ["A", "B"]})
        ]
        
        presentation = generator.generate_deck(intents)
        assert len(presentation.slides) == 3
    
    def test_adapt_element_count_add_steps(self):
        """Test adapting element count by adding steps."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.PROCESS_FLOW,
            content={"steps": ["Step 1", "Step 2"]}
        )
        
        generator.adapt_element_count(intent, 5)
        steps = intent.get_process_steps()
        assert len(steps) == 5
        assert steps[4].label == "Step 5"
    
    def test_adapt_element_count_remove_steps(self):
        """Test adapting element count by removing steps."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.PROCESS_FLOW,
            content={"steps": ["Step 1", "Step 2", "Step 3", "Step 4", "Step 5"]}
        )
        
        generator.adapt_element_count(intent, 3)
        steps = intent.get_process_steps()
        assert len(steps) == 3
        assert steps[2].label == "Step 3"
    
    def test_adapt_element_count_no_change(self):
        """Test adapting element count when already matching."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.PROCESS_FLOW,
            content={"steps": ["Step 1", "Step 2", "Step 3"]}
        )
        
        generator.adapt_element_count(intent, 3)
        steps = intent.get_process_steps()
        assert len(steps) == 3
    
    def test_adapt_element_count_non_process_flow(self):
        """Test adapting element count for non-process flow slide."""
        generator = SlideGenerator()
        
        intent = GenerationIntent(
            slide_type=SlideType.TITLE_SLIDE,
            content={"title": "Test"}
        )
        
        generator.adapt_element_count(intent, 5)
        assert intent.content["title"] == "Test"
    
    def test_empty_text_edge_cases(self):
        """Test that empty or blank text fields do not crash with IndexError."""
        generator = SlideGenerator()
        slide = generator.presentation.slides.add_slide(generator.presentation.slide_layouts[6])
        
        # Calling with empty strings should be completely safe
        generator._add_title(slide, "")
        generator._add_subtitle(slide, "")
        generator._add_bullet_points(slide, ["", "Valid point", ""])
        generator._add_body_text(slide, "")
        
        empty_intent = GenerationIntent(
            slide_type=SlideType.COMPARISON,
            content={"bullet_points": ["", ""]}
        )
        generator._add_comparison_layout(slide, empty_intent)
