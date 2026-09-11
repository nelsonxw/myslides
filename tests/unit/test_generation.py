"""
Unit tests for slide composition, validation, and builder.
"""
from pathlib import Path
import pytest
from pptx import Presentation

from myslides.generation.builder import build_pptx_slide
from myslides.generation.composer import (
    _build_fallback_blueprint,
    compose_slide_blueprint,
    validate_and_autofix_blueprint,
)
from myslides.generation.intent import parse_user_intent
from myslides.generation.spec import ShapeBlueprint, SlideBlueprint, SlideIntentSpec
from myslides.llm.dell_gateway import MockLLMProvider


def test_intent_parsing_and_fallback():
    mock_llm = MockLLMProvider()
    intent = parse_user_intent("Quarterly KPI update with 20% revenue growth", mock_llm)
    assert isinstance(intent, SlideIntentSpec)
    assert intent.headline != ""
    assert intent.archetype in ("kpi_summary", "general", "data_chart")


def test_blueprint_validation_and_bounds_clamping():
    # Blueprint with out-of-bounds coordinates
    bp = SlideBlueprint(
        title="Test Slide",
        subtitle="Test Subtitle",
        archetype="kpi_summary",
        shapes=[
            ShapeBlueprint(
                id="c1",
                type="card",
                left_in=0.1,  # too far left
                top_in=0.5,   # overlaps title
                width_in=15.0,# wider than 13.33in canvas
                height_in=10.0,
            )
        ],
    )
    fixed = validate_and_autofix_blueprint(bp)
    s = fixed.shapes[0]
    assert s.left_in >= 0.6
    assert s.top_in >= 1.4
    assert s.left_in + s.width_in <= 13.33
    assert s.top_in + s.height_in <= 7.5


def test_build_pptx_slide(tmp_path):
    bp = SlideBlueprint(
        title="Revenue & Performance Overview",
        subtitle="Executive Summary Q3 2026",
        archetype="kpi_summary",
        shapes=[
            ShapeBlueprint(
                id="c1",
                type="card",
                left_in=0.8,
                top_in=1.6,
                width_in=3.6,
                height_in=4.8,
                bg_color="#F0F0F0",
                title="Revenue Growth",
                stat_number="+24%",
                stat_label="Year over Year",
                items=["Exceeded Q3 revenue target by $4.2M", "Record operating margin in enterprise"],
            ),
            ShapeBlueprint(
                id="c2",
                type="card",
                left_in=4.8,
                top_in=1.6,
                width_in=3.6,
                height_in=4.8,
                bg_color="#FFFFFF",
                title="Cloud Adoption",
                stat_number="94%",
                stat_label="Customer Migration",
                items=["48 enterprise clients fully onboarded", "Zero downtime during transition"],
            ),
            ShapeBlueprint(
                id="c3",
                type="card",
                left_in=8.8,
                top_in=1.6,
                width_in=3.6,
                height_in=4.8,
                bg_color="#F0F0F0",
                title="Customer CSAT",
                stat_number="4.9 / 5",
                stat_label="Quarterly Score",
                items=["Top ranking across all regional accounts", "NPS reached an all-time peak of +72"],
            ),
        ],
    )

    out_file = tmp_path / "test_output.pptx"
    result_path = build_pptx_slide(bp, out_file)

    assert result_path.exists()
    assert result_path.stat().st_size > 1000

    # Verify presentation structure with python-pptx
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert len(slide.shapes) >= 3
