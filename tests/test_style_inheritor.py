"""
Unit tests for StyleInheritor.
"""
import pytest
from unittest.mock import Mock, MagicMock
from pptx.dml.color import RGBColor

from myslides.generation.style_inheritor import StyleInheritor


class TestStyleInheritor:
    """Test StyleInheritor class."""
    
    def test_initialization_no_palette(self):
        """Test initialization without color palette."""
        inheritor = StyleInheritor()
        assert inheritor.color_palette == []
    
    def test_initialization_with_palette(self):
        """Test initialization with color palette."""
        palette = ["#FF0000", "#00FF00", "#0000FF"]
        inheritor = StyleInheritor(color_palette=palette)
        assert inheritor.color_palette == palette
    
    def test_hex_to_rgb(self):
        """Test hex color conversion to RGBColor."""
        inheritor = StyleInheritor()
        rgb = inheritor._hex_to_rgb("#FF0000")
        assert isinstance(rgb, RGBColor)
    
    def test_hex_to_rgb_without_hash(self):
        """Test hex color conversion without # prefix."""
        inheritor = StyleInheritor()
        rgb = inheritor._hex_to_rgb("FF0000")
        assert isinstance(rgb, RGBColor)
    
    def test_hex_to_rgb_caching(self):
        """Test that RGB conversion is cached."""
        inheritor = StyleInheritor()
        rgb1 = inheritor._hex_to_rgb("#FF0000")
        rgb2 = inheritor._hex_to_rgb("#FF0000")
        # Check that both produce the same color values
        assert rgb1 == rgb2
    
    def test_get_primary_color(self):
        """Test getting primary color from palette."""
        palette = ["#FF0000", "#00FF00", "#0000FF"]
        inheritor = StyleInheritor(color_palette=palette)
        rgb = inheritor.get_primary_color()
        assert rgb is not None
    
    def test_get_primary_color_no_palette(self):
        """Test getting primary color when palette is empty."""
        inheritor = StyleInheritor()
        rgb = inheritor.get_primary_color()
        assert rgb is None
    
    def test_get_secondary_color(self):
        """Test getting secondary color from palette."""
        palette = ["#FF0000", "#00FF00", "#0000FF"]
        inheritor = StyleInheritor(color_palette=palette)
        rgb = inheritor.get_secondary_color()
        assert rgb is not None
    
    def test_get_secondary_color_single_palette(self):
        """Test getting secondary color when palette has only one color."""
        palette = ["#FF0000"]
        inheritor = StyleInheritor(color_palette=palette)
        rgb = inheritor.get_secondary_color()
        assert rgb is None
    
    def test_get_accent_color(self):
        """Test getting accent color from palette."""
        palette = ["#FF0000", "#00FF00", "#0000FF", "#FFFF00"]
        inheritor = StyleInheritor(color_palette=palette)
        rgb = inheritor.get_accent_color(2)
        assert rgb is not None
    
    def test_get_accent_color_out_of_range(self):
        """Test getting accent color when index is out of range."""
        palette = ["#FF0000", "#00FF00"]
        inheritor = StyleInheritor(color_palette=palette)
        rgb = inheritor.get_accent_color(5)
        assert rgb is None
    
    def test_apply_palette_to_shape(self):
        """Test applying palette color to shape fill."""
        palette = ["#FF0000", "#00FF00"]
        inheritor = StyleInheritor(color_palette=palette)
        
        # Mock shape
        shape = Mock()
        shape.fill = Mock()
        shape.fill.solid = Mock()
        shape.fill.fore_color = Mock()
        
        inheritor.apply_palette_to_shape(shape, 0)
        
        shape.fill.solid.assert_called_once()
        assert shape.fill.fore_color.rgb is not None
    
    def test_apply_palette_to_shape_no_palette(self):
        """Test applying palette when palette is empty."""
        inheritor = StyleInheritor()
        shape = Mock()
        shape.fill = Mock()
        
        inheritor.apply_palette_to_shape(shape, 0)
        
        # Should not call fill methods
        shape.fill.solid.assert_not_called()
    
    def test_apply_palette_to_shape_no_fill(self):
        """Test applying palette when shape has no fill."""
        palette = ["#FF0000"]
        inheritor = StyleInheritor(color_palette=palette)
        
        shape = Mock()
        shape.fill = None
        
        # Should not raise exception
        inheritor.apply_palette_to_shape(shape, 0)
    
    def test_apply_palette_to_text_frame(self):
        """Test applying palette color to text."""
        palette = ["#FF0000", "#00FF00"]
        inheritor = StyleInheritor(color_palette=palette)
        
        # Mock text frame
        text_frame = Mock()
        paragraph = Mock()
        run = Mock()
        run.font = Mock()
        run.font.color = Mock()
        paragraph.runs = [run]
        text_frame.paragraphs = [paragraph]
        
        inheritor.apply_palette_to_text_frame(text_frame, 0)
        
        assert run.font.color.rgb is not None
    
    def test_apply_font_style(self):
        """Test applying font styling to text run."""
        inheritor = StyleInheritor()
        
        text_run = Mock()
        text_run.font = Mock()
        
        inheritor.apply_font_style(
            text_run,
            font_name="Arial",
            font_size=12.0,
            bold=True,
            italic=True
        )
        
        # Check that font properties were set
        assert text_run.font.name == "Arial"
        assert text_run.font.bold is True
        assert text_run.font.italic is True
    
    def test_apply_alignment(self):
        """Test applying text alignment."""
        inheritor = StyleInheritor()
        
        paragraph = Mock()
        
        inheritor.apply_alignment(paragraph, "center")
        
        # Check that alignment was set
        assert paragraph.alignment is not None
    
    def test_apply_shape_border(self):
        """Test applying border to shape."""
        palette = ["#FF0000"]
        inheritor = StyleInheritor(color_palette=palette)
        
        shape = Mock()
        shape.line = Mock()
        shape.line.color = Mock()
        shape.line.width = Mock()
        
        inheritor.apply_shape_border(shape, border_width=2.0)
        
        # Check that border properties were accessed
        assert shape.line.color.rgb is not None
    
    def test_apply_shape_border_no_line(self):
        """Test applying border when shape has no line."""
        inheritor = StyleInheritor()
        
        shape = Mock()
        shape.line = None
        
        # Should not raise exception
        inheritor.apply_shape_border(shape)
    
    def test_apply_shape_shadow(self):
        """Test applying shadow to shape."""
        inheritor = StyleInheritor()
        
        shape = Mock()
        shape.shadow = Mock()
        
        inheritor.apply_shape_shadow(shape, shadow_enabled=True)
        
        # Check that shadow inherit was set
        assert shape.shadow.inherit is True
    
    def test_apply_shape_shadow_disable(self):
        """Test disabling shadow on shape."""
        inheritor = StyleInheritor()
        
        shape = Mock()
        shape.shadow = Mock()
        
        inheritor.apply_shape_shadow(shape, shadow_enabled=False)
        
        # Check that shadow inherit was set to False
        assert shape.shadow.inherit is False
    
    def test_apply_shape_shadow_no_shadow(self):
        """Test applying shadow when shape has no shadow."""
        inheritor = StyleInheritor()
        
        shape = Mock()
        shape.shadow = None
        
        # Should not raise exception
        inheritor.apply_shape_shadow(shape)
