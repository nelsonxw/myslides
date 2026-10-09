"""
Tests for SlideRedesignService and improve-slide endpoint.
"""
from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from myslides.api.main import app
from myslides.generation.slide_redesign_service import (
    SlideRedesignService,
    SlideAnalysisResult,
)
from myslides.ingestion.pptx_parser import (
    SlideInfo,
    ShapeInfo,
    PositionInfo,
    TextContent,
    FontInfo,
    SlideLayoutType,
)


@pytest.fixture
def client():
    """Create test client for FastAPI app."""
    return TestClient(app)


@pytest.fixture
def sample_slide_info():
    """Create sample parsed slide info with text-heavy bullets."""
    font = FontInfo(name="Calibri", size=14.0, bold=False, italic=False, underline=False, color=None)
    title_content = TextContent(text="Q4 Product Strategy", font=font, paragraphs=["Q4 Product Strategy"])
    body_content = TextContent(
        text="Phase 1: Design sprint\nPhase 2: Alpha rollout\nPhase 3: Global GA release",
        font=font,
        paragraphs=[
            "Phase 1: Design sprint",
            "Phase 2: Alpha rollout",
            "Phase 3: Global GA release",
        ],
    )

    pos = PositionInfo(x=100, y=100, width=500, height=300)
    shapes = [
        ShapeInfo(
            shape_id=1,
            shape_type="text_box",
            name="Title",
            position=pos,
            fill_color=None,
            line_color=None,
            text_content=title_content,
            is_grouped=False,
            group_id=None,
        ),
        ShapeInfo(
            shape_id=2,
            shape_type="text_box",
            name="Body",
            position=pos,
            fill_color=None,
            line_color=None,
            text_content=body_content,
            is_grouped=False,
            group_id=None,
        ),
    ]

    return SlideInfo(
        slide_index=0,
        layout_type=SlideLayoutType.CONTENT_SLIDE,
        width=9144000,
        height=5143500,
        shapes=shapes,
        charts=[],
        tables=[],
        images=[],
        text_content="Q4 Product Strategy\nPhase 1: Design sprint\nPhase 2: Alpha rollout\nPhase 3: Global GA release",
        color_palette=["#003366"],
        complexity_score=5,
    )


class TestSlideRedesignService:
    """Unit tests for SlideRedesignService."""

    def test_extract_content_elements(self, sample_slide_info):
        """Test extraction of title and bullet items from slide info."""
        service = SlideRedesignService()
        title, subtitle, items = service._extract_content_elements(sample_slide_info)

        assert title == "Q4 Product Strategy"
        assert len(items) == 3
        assert items[0] == "Phase 1: Design sprint"
        assert items[1] == "Phase 2: Alpha rollout"
        assert items[2] == "Phase 3: Global GA release"

    def test_discover_learned_layouts(self, sample_slide_info):
        """Test that layout recommendations are discovered directly from curated slides in training."""
        mock_matcher = MagicMock()
        mock_matcher.semantic_collection.query.return_value = {
            "ids": [["tmpl_1", "tmpl_2"]],
            "distances": [[100.0, 200.0]],
        }
        mock_db = MagicMock()
        mock_template_1 = MagicMock()
        mock_template_1.classification = "process_flow"
        mock_template_1.to_dict.return_value = {"id": "tmpl_1", "classification": "process_flow"}
        mock_template_2 = MagicMock()
        mock_template_2.classification = "comparison"
        mock_template_2.to_dict.return_value = {"id": "tmpl_2", "classification": "comparison"}
        mock_db.get_template_by_id.side_effect = lambda tid: mock_template_1 if tid == "tmpl_1" else mock_template_2

        service = SlideRedesignService(template_matcher=mock_matcher, database_manager=mock_db)
        templates, categories = service._discover_learned_layouts(sample_slide_info)

        assert len(templates) == 2
        assert categories == ["process_flow", "comparison"]
        assert templates[0]["template_id"] == "tmpl_1"
        assert templates[0]["similarity_score"] > 0

    def test_analyze_slide_info(self, sample_slide_info):
        """Test end-to-end analysis of SlideInfo into suggestions."""
        mock_matcher = MagicMock()
        mock_matcher.find_similar_templates.return_value = []

        service = SlideRedesignService(template_matcher=mock_matcher)
        result = service.analyze_slide_info(sample_slide_info)

        assert isinstance(result, SlideAnalysisResult)
        assert result.extracted_title == "Q4 Product Strategy"
        assert len(result.extracted_items) == 3
        assert len(result.recommended_slide_types) > 0
        assert "parsed_intent" in result.to_dict()


class TestImproveSlideEndpoint:
    """Tests for POST /api/generate/improve-slide."""

    def test_improve_slide_endpoint(self, client, tmp_path):
        """Test uploading a real minimal PPTX slide for improvement."""
        # Create a valid PPTX slide file
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        tx_box = slide.shapes.add_textbox(100, 100, 400, 200)
        tf = tx_box.text_frame
        tf.text = "Existing Architecture\nStep 1: Ingest data\nStep 2: Vector search\nStep 3: Generate deck"

        pptx_file = tmp_path / "test_slide.pptx"
        prs.save(str(pptx_file))

        with open(pptx_file, "rb") as f:
            response = client.post(
                "/api/generate/improve-slide",
                files={"file": ("test_slide.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
                data={"prompt": "Convert to modern process flow"},
            )

        assert response.status_code == 200
        data = response.json()
        assert "original_classification" in data
        assert "extracted_title" in data
        assert "recommended_slide_types" in data
        assert "suggested_prompt" in data
        assert "parsed_intent" in data
