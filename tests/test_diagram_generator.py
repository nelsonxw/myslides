"""
Unit tests for DiagramGenerator.
"""
import pytest
from pptx import Presentation

from myslides.generation.diagram_generator import DiagramGenerator
from myslides.generation.generation_models import ProcessStep, TimelineEvent
from myslides.generation.style_inheritor import StyleInheritor


class TestDiagramGenerator:
    """Test DiagramGenerator class."""
    
    def test_initialization(self):
        """Test DiagramGenerator initialization."""
        presentation = Presentation()
        generator = DiagramGenerator(presentation)
        assert generator.presentation is presentation
        assert generator.style_inheritor is not None
    
    def test_initialization_with_style_inheritor(self):
        """Test DiagramGenerator initialization with custom StyleInheritor."""
        presentation = Presentation()
        style_inheritor = StyleInheritor(color_palette=["#FF0000"])
        generator = DiagramGenerator(presentation, style_inheritor)
        assert generator.style_inheritor is style_inheritor
    
    def test_create_process_flow_horizontal(self):
        """Test creating horizontal process flow with real slide."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        style_inheritor = StyleInheritor(color_palette=["#1F4E79", "#2CA02C", "#D62728"])
        generator = DiagramGenerator(presentation, style_inheritor)
        
        steps = [
            ProcessStep(label="Step 1", number=1),
            ProcessStep(label="Step 2", number=2),
            ProcessStep(label="Step 3", number=3)
        ]
        
        generator.create_process_flow(slide, steps, horizontal=True)
        # 3 step shapes + 2 arrow shapes = 5 shapes
        assert len(slide.shapes) == 5
    
    def test_create_process_flow_vertical(self):
        """Test creating vertical process flow with real slide."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        style_inheritor = StyleInheritor(color_palette=["#1F4E79", "#2CA02C"])
        generator = DiagramGenerator(presentation, style_inheritor)
        
        steps = [
            ProcessStep(label="Step 1", number=1),
            ProcessStep(label="Step 2", number=2)
        ]
        
        generator.create_process_flow(slide, steps, horizontal=False)
        # 2 step shapes + 1 arrow shape = 3 shapes
        assert len(slide.shapes) == 3
    
    def test_create_process_flow_empty(self):
        """Test creating process flow with no steps."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)
        
        generator.create_process_flow(slide, [])
        assert len(slide.shapes) == 0
    
    def test_create_process_flow_with_description(self):
        """Test creating process flow with step descriptions."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        style_inheritor = StyleInheritor(color_palette=["#1F4E79"])
        generator = DiagramGenerator(presentation, style_inheritor)
        
        steps = [
            ProcessStep(label="Step 1", description="First step description", number=1)
        ]
        
        generator.create_process_flow(slide, steps, horizontal=True)
        assert len(slide.shapes) == 1
        shape = slide.shapes[0]
        assert "Step 1" in shape.text_frame.text
        assert "First step description" in shape.text_frame.text
    
    def test_create_timeline_horizontal(self):
        """Test creating horizontal timeline on real slide."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        style_inheritor = StyleInheritor(color_palette=["#1F4E79", "#2CA02C"])
        generator = DiagramGenerator(presentation, style_inheritor)
        
        events = [
            TimelineEvent(date="2024-01-01", title="Milestone 1"),
            TimelineEvent(date="2024-06-01", title="Milestone 2")
        ]
        
        generator.create_timeline(slide, events, horizontal=True)
        # 1 connector line + 2 oval markers + 2 date textboxes + 2 title textboxes = 7 shapes
        assert len(slide.shapes) == 7
    
    def test_create_timeline_vertical(self):
        """Test creating vertical timeline on real slide."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        style_inheritor = StyleInheritor(color_palette=["#1F4E79"])
        generator = DiagramGenerator(presentation, style_inheritor)
        
        events = [
            TimelineEvent(date="2024-01-01", title="Milestone 1")
        ]
        
        generator.create_timeline(slide, events, horizontal=False)
        # 1 connector line + 1 oval marker + 1 date textbox + 1 title textbox = 4 shapes
        assert len(slide.shapes) == 4
    
    def test_create_timeline_empty(self):
        """Test creating timeline with no events."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)
        
        generator.create_timeline(slide, [])
        assert len(slide.shapes) == 0
    
    def test_add_shape_types(self):
        """Test adding various shape types."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)
        
        shape_rect = generator._add_shape(slide, "rectangle", 100, 100, 200, 150)
        shape_chevron = generator._add_shape(slide, "chevron", 300, 100, 200, 150)
        shape_oval = generator._add_shape(slide, "oval", 500, 100, 200, 150)
        shape_diamond = generator._add_shape(slide, "diamond", 700, 100, 200, 150)
        shape_default = generator._add_shape(slide, "unknown_type", 900, 100, 200, 150)
        
        assert len(slide.shapes) == 5
        assert shape_rect is not None
        assert shape_chevron is not None
        assert shape_oval is not None
        assert shape_diamond is not None
        assert shape_default is not None
