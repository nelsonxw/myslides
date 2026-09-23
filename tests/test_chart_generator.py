"""
Unit tests for ChartGenerator.
"""
import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from myslides.generation.chart_generator import ChartGenerator
from myslides.generation.generation_models import ChartData, ChartSubType


class TestChartGenerator:
    """Test ChartGenerator class."""
    
    def test_initialization(self):
        """Test ChartGenerator initialization."""
        presentation = Presentation()
        generator = ChartGenerator(presentation)
        assert generator.presentation is presentation
    
    def test_get_pptx_chart_type_bar_vertical(self):
        """Test mapping bar vertical to pptx chart type."""
        presentation = Presentation()
        generator = ChartGenerator(presentation)
        
        chart_type = generator._get_pptx_chart_type(ChartSubType.BAR_VERTICAL)
        assert chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    
    def test_get_pptx_chart_type_bar_horizontal(self):
        """Test mapping bar horizontal to pptx chart type."""
        presentation = Presentation()
        generator = ChartGenerator(presentation)
        
        chart_type = generator._get_pptx_chart_type(ChartSubType.BAR_HORIZONTAL)
        assert chart_type == XL_CHART_TYPE.BAR_CLUSTERED
    
    def test_get_pptx_chart_type_pie(self):
        """Test mapping pie to pptx chart type."""
        presentation = Presentation()
        generator = ChartGenerator(presentation)
        
        chart_type = generator._get_pptx_chart_type(ChartSubType.PIE)
        assert chart_type == XL_CHART_TYPE.PIE
    
    def test_get_pptx_chart_type_line(self):
        """Test mapping line to pptx chart type."""
        presentation = Presentation()
        generator = ChartGenerator(presentation)
        
        chart_type = generator._get_pptx_chart_type(ChartSubType.LINE_SINGLE)
        assert chart_type == XL_CHART_TYPE.LINE
    
    def test_get_pptx_chart_type_default(self):
        """Test default chart type mapping."""
        presentation = Presentation()
        generator = ChartGenerator(presentation)
        
        chart_type = generator._get_pptx_chart_type(ChartSubType.WATERFALL)
        assert chart_type == XL_CHART_TYPE.COLUMN_STACKED
    
    def test_add_chart_to_slide_real(self):
        """Test adding a real chart to a presentation slide."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)
        
        chart_data = ChartData(
            chart_type=ChartSubType.BAR_VERTICAL,
            title="Revenue Q1-Q3",
            categories=["Q1", "Q2", "Q3"],
            series_data={"Revenue": [100.0, 150.0, 200.0]},
            has_legend=True,
            show_data_labels=True
        )
        
        generator.add_chart_to_slide(slide, chart_data)
        assert len(slide.shapes) == 1
        chart = slide.shapes[0].chart
        assert chart.has_title is True
        assert chart.chart_title.text_frame.text == "Revenue Q1-Q3"
        assert chart.has_legend is True
    
    def test_add_chart_no_legend(self):
        """Test adding chart with legend disabled."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)
        
        chart_data = ChartData(
            chart_type=ChartSubType.BAR_VERTICAL,
            title="No Legend Chart",
            categories=["A", "B"],
            series_data={"Series": [10.0, 20.0]},
            has_legend=False
        )
        
        generator.add_chart_to_slide(slide, chart_data)
        chart = slide.shapes[0].chart
        assert chart.has_legend is False
    
    def test_create_bar_chart(self):
        """Test convenience method for creating bar chart."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)
        
        generator.create_bar_chart(
            slide,
            title="Sales",
            categories=["Q1", "Q2", "Q3"],
            series_data={"Sales": [100.0, 150.0, 200.0]},
            horizontal=False,
            stacked=False
        )
        assert len(slide.shapes) == 1
    
    def test_create_bar_chart_horizontal(self):
        """Test creating horizontal bar chart."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)
        
        generator.create_bar_chart(
            slide,
            title="Sales Horizontal",
            categories=["Q1", "Q2"],
            series_data={"Sales": [100.0, 150.0]},
            horizontal=True
        )
        assert len(slide.shapes) == 1
    
    def test_create_line_chart(self):
        """Test convenience method for creating line chart."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)
        
        generator.create_line_chart(
            slide,
            title="Monthly Trend",
            categories=["Jan", "Feb", "Mar"],
            series_data={"Active Users": [10.0, 20.0, 30.0]}
        )
        assert len(slide.shapes) == 1
    
    def test_create_pie_chart(self):
        """Test convenience method for creating pie chart."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)
        
        generator.create_pie_chart(
            slide,
            title="Market Share",
            categories=["Product A", "Product B", "Product C"],
            values=[40.0, 35.0, 25.0]
        )
        assert len(slide.shapes) == 1
