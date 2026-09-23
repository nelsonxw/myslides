"""
Tests for advanced diagram generation (FR-4.3 Phase 2).
Tests cycles, hierarchies, pyramids, matrices, and Venn diagrams.
"""
import pytest
from pptx import Presentation
from myslides.generation.diagram_generator import DiagramGenerator
from myslides.generation.style_inheritor import StyleInheritor
from myslides.generation.generation_models import ProcessStep


class TestAdvancedDiagrams:
    """Test advanced diagram types."""

    def test_cycle_diagram(self):
        """Test circular cycle diagram creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)

        steps = [
            ProcessStep(number=1, label="Plan", description="Planning phase"),
            ProcessStep(number=2, label="Execute", description="Execution phase"),
            ProcessStep(number=3, label="Review", description="Review phase")
        ]

        generator.create_cycle_diagram(slide, steps)

        assert len(slide.shapes) > 0

    def test_hierarchy_diagram(self):
        """Test hierarchy/org chart diagram creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)

        hierarchy = {
            "root": {"label": "CEO"},
            "children": [
                {"label": "VP Engineering"},
                {"label": "VP Marketing"},
                {"label": "VP Sales"}
            ]
        }

        generator.create_hierarchy_diagram(slide, hierarchy)

        assert len(slide.shapes) > 0

    def test_pyramid_diagram(self):
        """Test pyramid/funnel diagram creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)

        levels = [
            {"label": "Level 1", "value": 100},
            {"label": "Level 2", "value": 75},
            {"label": "Level 3", "value": 50}
        ]

        generator.create_pyramid_diagram(slide, levels)

        assert len(slide.shapes) > 0

    def test_matrix_diagram(self):
        """Test 2x2 matrix/quadrant diagram creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)

        quadrants = [
            {"label": "Q1", "description": "High growth"},
            {"label": "Q2", "description": "High share"},
            {"label": "Q3", "description": "Low share"},
            {"label": "Q4", "description": "Low growth"}
        ]

        generator.create_matrix_diagram(slide, quadrants)

        assert len(slide.shapes) > 0

    def test_venn_diagram(self):
        """Test Venn diagram creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = DiagramGenerator(presentation)

        sets = [
            {"label": "Set A"},
            {"label": "Set B"},
            {"label": "Set C"}
        ]

        generator.create_venn_diagram(slide, sets)

        assert len(slide.shapes) > 0
