"""
Tests for slide preview renderer and PDF exporter (FR-4.5 Phase 2).
"""
import pytest
from myslides.generation.slide_preview_renderer import SlidePreviewRenderer
from myslides.generation.pdf_exporter import PDFExporter
import tempfile
from pathlib import Path


class TestSlidePreviewRenderer:
    """Test slide preview rendering."""

    def test_get_slide_metadata(self):
        """Test getting slide metadata."""
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        # Add a text box
        slide.shapes.add_textbox(1, 1, 2, 1)

        temp_path = tempfile.mktemp(suffix='.pptx')
        presentation.save(temp_path)

        metadata = SlidePreviewRenderer.get_slide_metadata(temp_path, 0)

        assert metadata["slide_index"] == 0
        assert metadata["shape_count"] >= 1
        assert "has_chart" in metadata
        assert "has_table" in metadata

        Path(temp_path).unlink(missing_ok=True)

    def test_render_slide_to_image_placeholder(self):
        """Test render_slide_to_image returns placeholder path."""
        temp_path = tempfile.mktemp(suffix='.pptx')

        result = SlidePreviewRenderer.render_slide_to_image(temp_path, 0)

        # For MVP, returns the PPTX path as placeholder
        assert result == temp_path

    def test_render_all_slides_placeholder(self):
        """Test render_all_slides returns placeholder paths."""
        from pptx import Presentation

        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        temp_path = tempfile.mktemp(suffix='.pptx')
        presentation.save(temp_path)

        result = SlidePreviewRenderer.render_all_slides(temp_path)

        assert len(result) >= 1

        Path(temp_path).unlink(missing_ok=True)


class TestPDFExporter:
    """Test PDF export functionality."""

    def test_export_to_pdf_placeholder(self):
        """Test export_to_pdf returns placeholder path."""
        temp_path = tempfile.mktemp(suffix='.pptx')

        result = PDFExporter.export_to_pdf(temp_path)

        # For MVP, returns the PPTX path as placeholder
        assert result == temp_path

    def test_export_deck_to_pdf_placeholder(self):
        """Test export_deck_to_pdf returns placeholder path."""
        temp_path = tempfile.mktemp(suffix='.pptx')

        result = PDFExporter.export_deck_to_pdf(temp_path, [0, 1])

        # For MVP, returns the PPTX path as placeholder
        assert result == temp_path
