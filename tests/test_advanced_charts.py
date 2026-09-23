"""
Tests for advanced chart generation (FR-4.2 Phase 2).
Tests combo, waterfall, and stacked area charts.
"""
import pytest
from pptx import Presentation
from myslides.generation.chart_generator import ChartGenerator
from myslides.generation.generation_models import ChartSubType


class TestAdvancedCharts:
    """Test advanced chart types."""

    def test_create_combo_chart(self):
        """Test combo chart (bar + line) creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)

        categories = ["Q1", "Q2", "Q3", "Q4"]
        primary_series = {"Revenue": [100, 120, 140, 160]}
        secondary_series = {"Growth": [5, 10, 8, 12]}

        generator.create_combo_chart(slide, "Revenue & Growth", categories, primary_series, secondary_series)

        assert len(slide.shapes) > 0
        chart = slide.shapes[0].chart
        assert chart is not None

    def test_create_waterfall_chart(self):
        """Test waterfall chart creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)

        categories = ["Start", "+ Change", "- Variance", "End"]
        values = [100, 20, -15, 105]

        generator.create_waterfall_chart(slide, "Variance Analysis", categories, values)

        assert len(slide.shapes) > 0
        chart = slide.shapes[0].chart
        assert chart is not None

    def test_create_stacked_area_chart(self):
        """Test stacked area chart creation."""
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        generator = ChartGenerator(presentation)

        categories = ["Q1", "Q2", "Q3", "Q4"]
        series_data = {
            "Product A": [30, 35, 40, 45],
            "Product B": [20, 25, 30, 35],
            "Product C": [10, 15, 20, 25]
        }

        generator.create_stacked_area_chart(slide, "Product Mix", categories, series_data)

        assert len(slide.shapes) > 0
        chart = slide.shapes[0].chart
        assert chart is not None

    def test_area_stacked_chart_type_mapping(self):
        """Test AREA_STACKED chart type mapping."""
        from myslides.generation.chart_generator import ChartGenerator
        from pptx.enum.chart import XL_CHART_TYPE

        generator = ChartGenerator(Presentation())
        chart_type = generator._get_pptx_chart_type(ChartSubType.AREA_STACKED)

        assert chart_type == XL_CHART_TYPE.AREA_STACKED
