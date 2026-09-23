"""
Tests for PPTX Parser module.
"""
import pytest
from pathlib import Path
from unittest.mock import Mock, MagicMock
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from myslides.ingestion.pptx_parser import (
    PPTXParser, 
    SlideInfo, 
    SlideLayoutType, 
    ChartType,
    ShapeInfo,
    ColorInfo,
    FontInfo,
    PositionInfo
)


class TestColorInfo:
    """Tests for ColorInfo class."""
    
    def test_from_color_format_with_valid_color(self):
        """Test ColorInfo creation from valid color format."""
        mock_color = Mock()
        mock_color.type = 1  # RGB color
        mock_color.rgb = 0xFF0000  # Red
        
        color_info = ColorInfo.from_color_format(mock_color)
        
        assert color_info is not None
        assert color_info.hex_value == "#ff0000"
        assert color_info.rgb == (255, 0, 0)
        assert 0 <= color_info.brightness <= 1
    
    def test_from_color_format_with_no_color(self):
        """Test ColorInfo creation from no color."""
        mock_color = Mock()
        mock_color.type = 0  # No color
        
        color_info = ColorInfo.from_color_format(mock_color)
        
        assert color_info is None
    
    def test_from_color_format_with_none(self):
        """Test ColorInfo creation from None."""
        color_info = ColorInfo.from_color_format(None)
        
        assert color_info is None


class TestFontInfo:
    """Tests for FontInfo class."""
    
    def test_from_font(self):
        """Test FontInfo creation from font object."""
        mock_font = Mock()
        mock_font.name = "Arial"
        mock_font.size = Mock()
        mock_font.size.pt = 14.0
        mock_font.bold = True
        mock_font.italic = False
        mock_font.underline = False
        mock_font.color = None
        
        font_info = FontInfo.from_font(mock_font)
        
        assert font_info.name == "Arial"
        assert font_info.size == 14.0
        assert font_info.bold is True
        assert font_info.italic is False
        assert font_info.underline is False


class TestPositionInfo:
    """Tests for PositionInfo class."""
    
    def test_from_shape(self):
        """Test PositionInfo creation from shape."""
        mock_shape = Mock()
        mock_shape.left = 100
        mock_shape.top = 200
        mock_shape.width = 300
        mock_shape.height = 400
        
        position_info = PositionInfo.from_shape(mock_shape)
        
        assert position_info.x == 100
        assert position_info.y == 200
        assert position_info.width == 300
        assert position_info.height == 400


class TestShapeInfo:
    """Tests for ShapeInfo class."""
    
    def test_from_shape_basic(self):
        """Test ShapeInfo creation from basic shape."""
        mock_shape = Mock()
        mock_shape.shape_id = 1
        mock_shape.shape_type = MSO_SHAPE_TYPE.RECTANGLE
        mock_shape.name = "Rectangle 1"
        mock_shape.left = 100
        mock_shape.top = 200
        mock_shape.width = 300
        mock_shape.height = 400
        mock_shape.fill.foreground_color = None
        mock_shape.line.color = None
        mock_shape.text_frame = None
        mock_shape.shape_type = MSO_SHAPE_TYPE.RECTANGLE
        
        shape_info = ShapeInfo.from_shape(mock_shape)
        
        assert shape_info.shape_id == 1
        assert shape_info.shape_type == str(MSO_SHAPE_TYPE.RECTANGLE)
        assert shape_info.name == "Rectangle 1"
        assert shape_info.is_grouped is False


class TestPPTXParser:
    """Tests for PPTXParser class."""
    
    @pytest.fixture
    def mock_presentation(self):
        """Create a mock presentation object."""
        mock_pres = Mock(spec=Presentation)
        mock_pres.slide_width = 9144000  # Standard 10" width in EMUs
        mock_pres.slide_height = 6858000  # Standard 7.5" height in EMUs
        return mock_pres
    
    @pytest.fixture
    def sample_pptx_file(self, tmp_path):
        """Create a sample PPTX file for testing."""
        # This would normally create a real PPTX file
        # For unit tests, we'll mock the file existence
        pptx_path = tmp_path / "test.pptx"
        pptx_path.touch()  # Create empty file
        return pptx_path
    
    def test_parser_initialization_with_valid_file(self, sample_pptx_file):
        """Test parser initialization with valid file."""
        # We need to mock the Presentation class since we don't have a real PPTX
        with pytest.mock.patch('myslides.ingestion.pptx_parser.Presentation') as mock_presentation_class:
            mock_pres = Mock()
            mock_pres.slides = []
            mock_presentation_class.return_value = mock_pres
            
            parser = PPTXParser(sample_pptx_file)
            
            assert parser.pptx_path == sample_pptx_file
            assert parser.presentation == mock_pres
    
    def test_parser_initialization_with_invalid_file(self):
        """Test parser initialization with invalid file."""
        with pytest.raises(FileNotFoundError):
            PPTXParser("nonexistent.pptx")
    
    def test_parse_single_slide(self, sample_pptx_file):
        """Test parsing single slide from individual slide file."""
        with pytest.mock.patch('myslides.ingestion.pptx_parser.Presentation') as mock_presentation_class:
            mock_pres = Mock()
            mock_pres.slides = [Mock()]  # Single slide
            mock_pres.slide_width = 9144000
            mock_pres.slide_height = 6858000
            mock_presentation_class.return_value = mock_pres
            
            parser = PPTXParser(sample_pptx_file)
            slide = parser.parse_slide()
            
            assert slide.slide_index == 0
            assert slide.width == 9144000
            assert slide.height == 6858000
    
    def test_get_presentation_metadata(self, sample_pptx_file):
        """Test getting individual slide file metadata."""
        with pytest.mock.patch('myslides.ingestion.pptx_parser.Presentation') as mock_presentation_class:
            mock_pres = Mock()
            mock_pres.slides = [Mock()]  # Single slide
            mock_pres.slide_width = 9144000
            mock_pres.slide_height = 6858000
            mock_presentation_class.return_value = mock_pres
            
            parser = PPTXParser(sample_pptx_file)
            metadata = parser.get_presentation_metadata()
            
            assert metadata["file_name"] == sample_pptx_file.name
            assert metadata["slide_count"] == 1  # Individual slides always have 1 slide
            assert metadata["width"] == 9144000
            assert metadata["height"] == 6858000
    
    def test_determine_layout_type_title_slide(self):
        """Test layout type determination for title slide."""
        parser = PPTXParser.__new__(PPTXParser)  # Create without __init__
        
        mock_slide = Mock()
        mock_slide.shapes = [Mock(), Mock()]
        
        # Setup first shape as title
        mock_title_shape = Mock()
        mock_title_shape.text_frame = Mock()
        mock_title_shape.text_frame.text = "Main Title"
        mock_title_shape.text_frame.paragraphs = [Mock()]
        mock_title_shape.text_frame.paragraphs[0].runs = [Mock()]
        mock_title_shape.text_frame.paragraphs[0].runs[0].font.size = Mock()
        mock_title_shape.text_frame.paragraphs[0].runs[0].font.size.pt = 36
        
        # Setup second shape as subtitle
        mock_subtitle_shape = Mock()
        mock_subtitle_shape.text_frame = Mock()
        mock_subtitle_shape.text_frame.text = "Subtitle"
        mock_subtitle_shape.text_frame.paragraphs = [Mock()]
        mock_subtitle_shape.text_frame.paragraphs[0].runs = [Mock()]
        mock_subtitle_shape.text_frame.paragraphs[0].runs[0].font.size = Mock()
        mock_subtitle_shape.text_frame.paragraphs[0].runs[0].font.size.pt = 18
        
        mock_slide.shapes = [mock_title_shape, mock_subtitle_shape]
        
        layout_type = parser._determine_layout_type(mock_slide)
        
        assert layout_type == SlideLayoutType.TITLE_SLIDE
    
    def test_calculate_complexity(self):
        """Test complexity score calculation."""
        parser = PPTXParser.__new__(PPTXParser)
        
        mock_shapes = [Mock() for _ in range(5)]
        for shape in mock_shapes:
            shape.is_grouped = False
            shape.text_content = Mock()
            shape.text_content.text = "Sample text"
        
        mock_charts = [Mock()]
        mock_tables = [Mock()]
        mock_images = [Mock()]
        
        complexity = parser._calculate_complexity(mock_shapes, mock_charts, mock_tables, mock_images)
        
        # Base score: 5 shapes * 1 + 1 chart * 5 + 1 table * 3 + 1 image * 2 = 15
        # Plus text length contribution
        assert complexity >= 15


class TestSlideInfo:
    """Tests for SlideInfo class."""
    
    def test_to_dict(self):
        """Test SlideInfo to dictionary conversion."""
        slide_info = SlideInfo(
            slide_index=0,
            layout_type=SlideLayoutType.TITLE_SLIDE,
            width=9144000,
            height=6858000,
            shapes=[],
            charts=[],
            tables=[],
            images=[],
            text_content="Test content",
            color_palette=["#FF0000", "#00FF00"],
            complexity_score=5
        )
        
        slide_dict = slide_info.to_dict()
        
        assert slide_dict["slide_index"] == 0
        assert slide_dict["layout_type"] == "title_slide"
        assert slide_dict["text_content"] == "Test content"
        assert slide_dict["color_palette"] == ["#FF0000", "#00FF00"]
        assert slide_dict["complexity_score"] == 5


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
