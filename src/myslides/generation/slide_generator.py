"""
Slide Generator for MySlides.
Handles template-based slide assembly (FR-4.1).
"""
import logging
from typing import Optional, Dict, Any, List
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from myslides.generation.generation_models import (
    GenerationIntent, GenerationConfig, SlideType, ProcessStep
)
from myslides.generation.chart_generator import ChartGenerator
from myslides.generation.diagram_generator import DiagramGenerator
from myslides.generation.style_inheritor import StyleInheritor

logger = logging.getLogger(__name__)


class SlideGenerator:
    """Generates slides from structured intent using templates."""
    
    def __init__(self, config: Optional[GenerationConfig] = None, presentation: Optional[Presentation] = None):
        """
        Initialize the slide generator.
        
        Args:
            config: GenerationConfig with template and style information
            presentation: Optional existing presentation; creates a new one if omitted
        """
        self.config = config or GenerationConfig()
        self._custom_presentation = presentation is not None
        self.presentation = presentation if presentation is not None else Presentation()
        
        # Initialize sub-components
        self.style_inheritor = StyleInheritor(
            color_palette=self.config.template_color_palette
        )
        self.chart_generator = ChartGenerator(self.presentation)
        self.diagram_generator = DiagramGenerator(
            self.presentation, 
            self.style_inheritor
        )

    def _reset_presentation_if_needed(self) -> None:
        """Reset internal presentation for clean single-slide generation if not using custom presentation."""
        if not self._custom_presentation and len(self.presentation.slides) > 0:
            self.presentation = Presentation()
            self.chart_generator = ChartGenerator(self.presentation)
            self.diagram_generator = DiagramGenerator(
                self.presentation,
                self.style_inheritor
            )
    
    def clone_and_populate_template(self, intent: GenerationIntent, template_path: str) -> Optional[Presentation]:
        """
        Clone a curated slide template from PPTX and populate user content into its structure,
        preserving all original visual design, shapes, icons, diagrams, and formatting.
        """
        try:
            prs = Presentation(template_path)
            if not prs.slides:
                return None
            slide = prs.slides[0]

            text_content = intent.get_text_content()
            title = text_content.title if text_content else intent.content.get("title", "")
            subtitle = text_content.subtitle if text_content else intent.content.get("subtitle")

            # Collect all text shapes recursively (including inside GroupShapes)
            text_shapes = []
            def collect(shapes):
                for s in shapes:
                    if s.has_text_frame and s.text_frame:
                        text_shapes.append(s)
                    if hasattr(s, "shapes"):
                        collect(s.shapes)
            collect(slide.shapes)

            # Sort by vertical position (top to bottom)
            text_shapes.sort(key=lambda s: (s.top, s.left))

            # Identify title shape
            title_shape = None
            for s in text_shapes:
                s_name = s.name.lower()
                if "title" in s_name or (s.top < prs.slide_height * 0.28 and len(s.text_frame.text.strip()) < 80):
                    title_shape = s
                    break

            if title_shape and title:
                tf = title_shape.text_frame
                if tf.paragraphs:
                    tf.paragraphs[0].text = title
                    if tf.paragraphs[0].runs:
                        tf.paragraphs[0].runs[0].font.bold = True

            # Identify subtitle shape
            subtitle_shape = None
            if subtitle:
                for s in text_shapes:
                    s_name = s.name.lower()
                    if s is not title_shape and ("subtitle" in s_name or (s.top < prs.slide_height * 0.38 and s.top > (title_shape.top if title_shape else 0))):
                        subtitle_shape = s
                        break
                if subtitle_shape:
                    if subtitle_shape.text_frame.paragraphs:
                        subtitle_shape.text_frame.paragraphs[0].text = subtitle

            # Prepare list of content items to place into template body shapes
            content_list = []
            if intent.content.get("steps"):
                for st in intent.content["steps"]:
                    if isinstance(st, dict):
                        if st.get("label"):
                            content_list.append(st["label"])
                        if st.get("description"):
                            content_list.append(st["description"])
                    else:
                        content_list.append(str(st))
            elif intent.content.get("events"):
                for ev in intent.content["events"]:
                    if isinstance(ev, dict):
                        date_str = ev.get("date", "")
                        title_str = ev.get("title", "")
                        desc_str = ev.get("description", "")
                        if date_str or title_str:
                            content_list.append(f"{date_str} - {title_str}" if date_str and title_str else (title_str or date_str))
                        if desc_str:
                            content_list.append(desc_str)
                    else:
                        content_list.append(str(ev))
            elif intent.content.get("items"):
                content_list = list(intent.content["items"])
            elif text_content and text_content.bullet_points:
                content_list = list(text_content.bullet_points)

            # Filter valid body shapes (skip slide numbers or single-letter shapes)
            body_shapes = [s for s in text_shapes if s is not title_shape and s is not subtitle_shape]
            valid_body_shapes = [s for s in body_shapes if not s.text_frame.text.strip().isdigit() and len(s.text_frame.text.strip()) > 2]

            # Populate body shapes
            for idx, s in enumerate(valid_body_shapes):
                tf = s.text_frame
                if idx < len(content_list):
                    if tf.paragraphs:
                        tf.paragraphs[0].text = content_list[idx]
                else:
                    # Clear sample placeholder text
                    sample_patterns = ["sample text", "insert your", "lorem", "dolor", "step ", "text here"]
                    curr_lower = tf.text.lower()
                    if any(p in curr_lower for p in sample_patterns):
                        tf.text = ""

            return prs
        except Exception as err:
            logger.warning("Could not clone template PPTX (%s), falling back to programmatic assembly", err)
            return None

    def generate_slide(self, intent: GenerationIntent) -> Presentation:
        """
        Generate a single slide from structured intent.
        If a curated template PPTX is configured, clones and populates it directly.
        
        Args:
            intent: GenerationIntent with slide specifications
        
        Returns:
            pptx Presentation object with the generated slide
        """
        # If a curated template file path is provided, clone and populate it directly
        if self.config.template_pptx_path and Path(self.config.template_pptx_path).exists():
            cloned_prs = self.clone_and_populate_template(intent, self.config.template_pptx_path)
            if cloned_prs is not None:
                return cloned_prs

        self._reset_presentation_if_needed()
        slide_layout = self._get_slide_layout(intent.slide_type)
        slide = self.presentation.slides.add_slide(slide_layout)
        
        if self.config.template_element_manifest:
            self._apply_template_background(slide)
        
        self._generate_content(slide, intent)
        return self.presentation
    
    def generate_deck(self, intents: List[GenerationIntent]) -> Presentation:
        """
        Generate a multi-slide presentation deck (FR-4.5).
        
        Args:
            intents: List of GenerationIntent objects, one per slide
        
        Returns:
            pptx Presentation object containing all generated slides
        """
        if not self._custom_presentation:
            self.presentation = Presentation()
            self.chart_generator = ChartGenerator(self.presentation)
            self.diagram_generator = DiagramGenerator(
                self.presentation,
                self.style_inheritor
            )
        for intent in intents:
            slide_layout = self._get_slide_layout(intent.slide_type)
            slide = self.presentation.slides.add_slide(slide_layout)
            if self.config.template_element_manifest:
                self._apply_template_background(slide)
            self._generate_content(slide, intent)
        return self.presentation
    
    def _get_slide_layout(self, slide_type: SlideType):
        """Get appropriate slide layout for the slide type."""
        return self.presentation.slide_layouts[6]  # Blank layout
    
    def _apply_template_background(self, slide) -> None:
        """Apply background from template if available."""
        if self.style_inheritor.color_palette:
            background = slide.background
            if background and hasattr(background, "fill") and background.fill:
                try:
                    background.fill.solid()
                    primary = self.style_inheritor.get_primary_color()
                    if primary:
                        background.fill.fore_color.rgb = primary
                except Exception:
                    pass
    
    def _generate_content(self, slide, intent: GenerationIntent) -> None:
        """Generate slide content based on intent."""
        text_content = intent.get_text_content()
        table_data = intent.content.get("table_data")
        
        if text_content and text_content.title:
            self._add_title(slide, text_content.title)
        
        if text_content and text_content.subtitle:
            self._add_subtitle(slide, text_content.subtitle)

        has_subtitle = bool(text_content and text_content.subtitle)
        chart_top = 2.4 if has_subtitle else 2.0
        chart_height = 4.4 if has_subtitle else 4.8
        
        if intent.slide_type in (SlideType.DATA_CHART, SlideType.DASHBOARD):
            chart_data = intent.get_chart_data()
            if chart_data:
                self.chart_generator.add_chart_to_slide(
                    slide, chart_data, top=chart_top, height=chart_height
                )
            elif intent.content.get("items"):
                self._add_bullet_points(slide, intent.content["items"])
        
        elif intent.slide_type == SlideType.PROCESS_FLOW:
            steps = intent.get_process_steps()
            if steps:
                self.diagram_generator.create_process_flow(slide, steps)
            elif intent.content.get("items"):
                self._add_bullet_points(slide, intent.content["items"])
        
        elif intent.slide_type == SlideType.TIMELINE:
            events = intent.get_timeline_events()
            if events:
                self.diagram_generator.create_timeline(slide, events)
            elif intent.content.get("items"):
                self._add_bullet_points(slide, intent.content["items"])
        
        elif intent.slide_type == SlideType.COMPARISON:
            self._add_comparison_layout(slide, intent)

        elif intent.slide_type == SlideType.TABLE_VIEW:
            if table_data:
                self._add_table(slide, table_data)
            elif intent.content.get("items"):
                self._add_bullet_points(slide, intent.content["items"])
        
        elif intent.slide_type == SlideType.TITLE_SLIDE:
            # If the slide actually has items or steps, don't leave them out!
            if intent.content.get("items"):
                self._add_bullet_points(slide, intent.content["items"])
            elif intent.content.get("steps"):
                steps = intent.get_process_steps()
                if steps:
                    self.diagram_generator.create_process_flow(slide, steps)
        
        else:
            if table_data:
                self._add_table(slide, table_data)
            else:
                bullets = (text_content.bullet_points if text_content else None) or intent.content.get("items")
                if bullets:
                    self._add_bullet_points(slide, bullets)
                elif text_content and text_content.body_text:
                    self._add_body_text(slide, text_content.body_text)
    
    def _add_table(self, slide, table_data: Dict[str, Any]) -> None:
        """Add formatted native table to slide with intelligent dynamic font scaling and fitting."""
        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])
        
        if not headers and not rows:
            return
        
        num_rows = len(rows) + (1 if headers else 0)
        num_cols = len(headers) if headers else (len(rows[0]) if rows else 0)
        
        if num_rows == 0 or num_cols == 0:
            return
        
        total_chars = sum(len(str(c)) for r in rows for c in r)
        max_cell_chars = max((len(str(c)) for r in rows for c in r), default=0)
        
        # Intelligent dynamic font scaling and vertical spacing
        if num_rows >= 8 or total_chars > 600 or max_cell_chars > 120:
            header_font_size = Pt(8.5)
            body_font_size = Pt(7.0)
            cell_margin_v = Inches(0.02)
            cell_margin_h = Inches(0.04)
            top = Inches(1.3)
            table_height = Inches(min(5.8, 0.42 * num_rows + 0.3))
        elif num_rows >= 5 or total_chars > 300:
            header_font_size = Pt(9.5)
            body_font_size = Pt(8.0)
            cell_margin_v = Inches(0.03)
            cell_margin_h = Inches(0.05)
            top = Inches(1.5)
            table_height = Inches(min(5.2, 0.45 * num_rows + 0.4))
        else:
            header_font_size = Pt(11.0)
            body_font_size = Pt(9.5)
            cell_margin_v = Inches(0.05)
            cell_margin_h = Inches(0.08)
            top = Inches(1.8)
            table_height = Inches(min(4.8, 0.5 * num_rows + 0.5))
        
        left = Inches(0.4)
        width = Inches(9.2)
        
        table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, table_height)
        table = table_shape.table
        
        # Calculate proportional column widths
        col_char_counts = [0] * num_cols
        if headers:
            for c_i, h in enumerate(headers[:num_cols]):
                col_char_counts[c_i] = max(col_char_counts[c_i], len(str(h)))
        for r in rows:
            for c_i, cell_val in enumerate(r[:num_cols]):
                col_char_counts[c_i] = max(col_char_counts[c_i], len(str(cell_val)))
        
        total_weight = sum(max(15, count) for count in col_char_counts)
        if total_weight > 0 and num_cols > 1:
            for c_i in range(num_cols):
                weight = max(15, col_char_counts[c_i]) / total_weight
                table.columns[c_i].width = Inches(max(1.2, 9.2 * weight))
        
        primary_color = self.style_inheritor.get_primary_color() or RGBColor(31, 78, 121)
        header_row_offset = 0
        
        if headers:
            header_row_offset = 1
            for col_idx, header_text in enumerate(headers[:num_cols]):
                cell = table.cell(0, col_idx)
                cell.text = str(header_text)
                cell.margin_top = cell_margin_v
                cell.margin_bottom = cell_margin_v
                cell.margin_left = cell_margin_h
                cell.margin_right = cell_margin_h
                cell.fill.solid()
                cell.fill.fore_color.rgb = primary_color
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = header_font_size
        
        for r_idx, row_values in enumerate(rows):
            target_row = r_idx + header_row_offset
            is_alt = (r_idx % 2 == 1)
            for c_idx, cell_value in enumerate(row_values[:num_cols]):
                cell = table.cell(target_row, c_idx)
                cell.text = str(cell_value)
                cell.margin_top = cell_margin_v
                cell.margin_bottom = cell_margin_v
                cell.margin_left = cell_margin_h
                cell.margin_right = cell_margin_h
                if is_alt:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(245, 247, 250)
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
                    for run in p.runs:
                        run.font.size = body_font_size

    def _add_title(self, slide, title: str) -> None:
        """Add title to slide with template styling."""
        title_box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(0.6),
            Inches(8.4),
            Inches(0.9)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.word_wrap = True
        
        primary_color = self.style_inheritor.get_primary_color()
        if title_frame.paragraphs:
            p = title_frame.paragraphs[0]
            if p.runs:
                self.style_inheritor.apply_font_style(
                    p.runs[0],
                    font_size=28.0,
                    bold=True
                )
                if primary_color:
                    p.runs[0].font.color.rgb = primary_color
            self.style_inheritor.apply_alignment(p, "center")
    
    def _add_subtitle(self, slide, subtitle: str) -> None:
        """Add subtitle to slide."""
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(1.55),
            Inches(8.4),
            Inches(0.4)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_frame.word_wrap = True
        
        if subtitle_frame.paragraphs:
            p = subtitle_frame.paragraphs[0]
            if p.runs:
                self.style_inheritor.apply_font_style(
                    p.runs[0],
                    font_size=13.0,
                    italic=True
                )
            self.style_inheritor.apply_alignment(p, "left")
    
    def _add_bullet_points(self, slide, bullet_points: List[str]) -> None:
        """Add bullet points to slide."""
        content_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(2.5),
            Inches(9.0),
            Inches(4.0)
        )
        text_frame = content_box.text_frame
        text_frame.word_wrap = True
        
        for i, point in enumerate(bullet_points):
            p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
            p.text = point
            p.level = 0
            p.space_after = Pt(12)
            if p.runs:
                self.style_inheritor.apply_font_style(
                    p.runs[0],
                    font_size=14.0
                )
    
    def _add_body_text(self, slide, body_text: str) -> None:
        """Add body text to slide."""
        content_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(2.5),
            Inches(9.0),
            Inches(4.0)
        )
        text_frame = content_box.text_frame
        text_frame.text = body_text
        text_frame.word_wrap = True
        
        if text_frame.paragraphs and text_frame.paragraphs[0].runs:
            self.style_inheritor.apply_font_style(
                text_frame.paragraphs[0].runs[0],
                font_size=14.0
            )
    
    def _add_comparison_layout(self, slide, intent: GenerationIntent) -> None:
        """Add comparison layout, supporting 2, 3, or 4 column modern cards."""
        columns = intent.content.get("columns")
        
        if columns and isinstance(columns, list) and len(columns) >= 2:
            num_cols = len(columns)
            col_width = Inches(min(4.2, (8.8 - (num_cols - 1) * 0.25) / num_cols))
            col_gap = Inches(0.25)
            top = Inches(1.4)
            card_height = Inches(5.6)
            palette = self.style_inheritor.color_palette or ["#1F4E79", "#2E75B6", "#269664", "#D97706"]
            
            for i, col in enumerate(columns):
                left = Inches(0.6) + i * (col_width + col_gap)
                
                # Header card
                hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_width, Inches(0.5))
                hdr.fill.solid()
                hex_c = palette[i % len(palette)]
                hdr.fill.fore_color.rgb = self.style_inheritor._hex_to_rgb(hex_c)
                hdr.line.fill.background()
                hp = hdr.text_frame.paragraphs[0]
                hp.text = col.get("title", f"Category {i+1}")
                hp.alignment = PP_ALIGN.CENTER
                if hp.runs:
                    hp.runs[0].font.bold = True
                    hp.runs[0].font.size = Pt(11)
                    hp.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                
                # Body card
                body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top + Inches(0.55), col_width, card_height - Inches(0.65))
                body.fill.solid()
                body.fill.fore_color.rgb = RGBColor(248, 250, 252)
                body.line.color.rgb = RGBColor(226, 232, 240)
                tf = body.text_frame
                tf.word_wrap = True
                tf.margin_left = Inches(0.12)
                tf.margin_right = Inches(0.12)
                tf.margin_top = Inches(0.12)
                
                points = col.get("points", [])
                for pt_idx, pt in enumerate(points):
                    p = tf.paragraphs[0] if pt_idx == 0 else tf.add_paragraph()
                    p.text = chr(8226) + " " + str(pt)
                    p.space_after = Pt(4)
                    if p.runs:
                        p.runs[0].font.size = Pt(8.5)
                        p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
            return

        # Fallback to standard 2-column layout (for unit test compatibility)
        text_content = intent.get_text_content()
        if not text_content or not text_content.bullet_points:
            return
        
        mid_point = len(text_content.bullet_points) // 2
        left_points = text_content.bullet_points[:mid_point]
        right_points = text_content.bullet_points[mid_point:]
        
        # Left column
        left_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(2.5),
            Inches(4.0),
            Inches(4.0)
        )
        left_frame = left_box.text_frame
        left_frame.word_wrap = True
        for i, point in enumerate(left_points):
            p = left_frame.paragraphs[0] if i == 0 else left_frame.add_paragraph()
            p.text = point
            p.level = 0
            if p.runs:
                self.style_inheritor.apply_font_style(p.runs[0], font_size=14.0)
        
        # Right column
        right_box = slide.shapes.add_textbox(
            Inches(5.0),
            Inches(2.5),
            Inches(4.0),
            Inches(4.0)
        )
        right_frame = right_box.text_frame
        right_frame.word_wrap = True
        for i, point in enumerate(right_points):
            p = right_frame.paragraphs[0] if i == 0 else right_frame.add_paragraph()
            p.text = point
            p.level = 0
            if p.runs:
                self.style_inheritor.apply_font_style(p.runs[0], font_size=14.0)
    
    def adapt_element_count(self, intent: GenerationIntent, target_count: int) -> None:
        """
        Adapt template to match required element count (FR-3.3).
        
        Args:
            intent: GenerationIntent with current content
            target_count: Target number of elements
        """
        if intent.slide_type == SlideType.PROCESS_FLOW:
            steps = intent.get_process_steps()
            current_count = len(steps)
            
            if current_count < target_count:
                for i in range(current_count, target_count):
                    steps.append(ProcessStep(label=f"Step {i + 1}", number=i + 1))
                intent.content["steps"] = [s.label for s in steps]
            elif current_count > target_count:
                intent.content["steps"] = [s.label for s in steps[:target_count]]
