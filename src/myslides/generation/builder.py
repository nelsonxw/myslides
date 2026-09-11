"""
Builder: converts a SlideBlueprint into a real, native .pptx PowerPoint slide.
Follows all brand rules (Dell template, 16:9 widescreen, custom shape cards, icons, charts).
"""
from __future__ import annotations

import os
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from myslides.config import settings
from myslides.generation.spec import ShapeBlueprint, SlideBlueprint


def hex_to_rgb(hex_str: str | None) -> RGBColor:
    """Convert hex string '#0672CB' to RGBColor."""
    if not hex_str:
        return RGBColor(0x1D, 0x2C, 0x3B)
    hex_clean = hex_str.lstrip("#")
    if len(hex_clean) != 6:
        return RGBColor(0x1D, 0x2C, 0x3B)
    try:
        r = int(hex_clean[0:2], 16)
        g = int(hex_clean[2:4], 16)
        b = int(hex_clean[4:6], 16)
        return RGBColor(r, g, b)
    except Exception:
        return RGBColor(0x1D, 0x2C, 0x3B)


def delete_all_existing_slides(prs: Presentation) -> None:
    """Clear default template slides so output contains only the newly generated slide."""
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        if rId is None:
            rId = prs.slides._sldIdLst[0].get("r:id")
        if rId:
            prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


def find_icon_file(icon_name: str) -> Path | None:
    """Look up matching icon in the Fluent icon library."""
    if not settings.icon_dir.exists():
        return None
    clean = icon_name.lower().replace(" ", "_").replace("-", "_")
    candidates = [
        settings.icon_dir / f"ic_fluent_{clean}_24_filled.png",
        settings.icon_dir / f"ic_fluent_{clean}_arrow_24_filled.png",
        settings.icon_dir / f"ic_fluent_{clean}_circle_24_filled.png",
    ]
    for c in candidates:
        if c.exists():
            return c
    # Fallback to general icon
    fallback = settings.icon_dir / "ic_fluent_target_arrow_24_filled.png"
    return fallback if fallback.exists() else None


def build_pptx_slide(blueprint: SlideBlueprint, output_file: Path | str) -> Path:
    """Construct PowerPoint presentation file from SlideBlueprint."""
    out_path = Path(output_file).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)

    template_path = settings.base_template_path
    if template_path.exists():
        prs = Presentation(str(template_path))
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    # Clean existing slides
    delete_all_existing_slides(prs)

    # Select layout (Layout 29: Title & Subtitle only)
    layout_idx = min(blueprint.layout_index, len(prs.slide_layouts) - 1)
    slide_layout = prs.slide_layouts[layout_idx]
    slide = prs.slides.add_slide(slide_layout)

    # Populate Header Placeholders (do NOT change font.name to preserve corporate branding)
    if len(slide.placeholders) > 0:
        slide.placeholders[0].text = blueprint.title
    if len(slide.placeholders) > 1 and blueprint.subtitle:
        slide.placeholders[1].text = blueprint.subtitle

    # Build custom shapes
    for shape_bp in blueprint.shapes:
        _render_shape_element(slide, shape_bp)

    prs.save(str(out_path))
    return out_path


def _render_shape_element(slide, shape_bp: ShapeBlueprint) -> None:
    """Render an individual visual component onto the slide."""
    left = Inches(shape_bp.left_in)
    top = Inches(shape_bp.top_in)
    width = Inches(shape_bp.width_in)
    height = Inches(shape_bp.height_in)

    if shape_bp.type in ("card", "badge"):
        # Card container box
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if shape_bp.type == "card" else MSO_SHAPE.RECTANGLE,
            left, top, width, height,
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = hex_to_rgb(shape_bp.bg_color or "#F0F0F0")
        
        # Border
        if shape_bp.border_color:
            shape.line.color.rgb = hex_to_rgb(shape_bp.border_color)
            shape.line.width = Pt(1.5)
        else:
            shape.line.color.rgb = hex_to_rgb("#C5D4E3")
            shape.line.width = Pt(1.0)

        # Content text frame inside card
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        # 1. Stat callout if present
        if shape_bp.stat_number:
            p_stat = tf.paragraphs[0]
            p_stat.text = shape_bp.stat_number
            p_stat.font.size = Pt(36)
            p_stat.font.bold = True
            p_stat.font.color.rgb = hex_to_rgb("#0672CB")  # Dell Blue
            
            if shape_bp.stat_label:
                p_lbl = tf.add_paragraph()
                p_lbl.text = shape_bp.stat_label.upper()
                p_lbl.font.size = Pt(10)
                p_lbl.font.bold = True
                p_lbl.font.color.rgb = hex_to_rgb("#40586D")  # Raven
                p_lbl.space_after = Pt(12)

        # 2. Card Title
        if shape_bp.title:
            p_title = tf.add_paragraph() if (shape_bp.stat_number or tf.text) else tf.paragraphs[0]
            p_title.text = shape_bp.title
            p_title.font.size = Pt(16)
            p_title.font.bold = True
            p_title.font.color.rgb = hex_to_rgb("#1D2C3B")
            p_title.space_after = Pt(8)

        # 3. Bullet list / items
        for item in shape_bp.items:
            p_item = tf.add_paragraph()
            p_item.text = f"• {item}"
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = hex_to_rgb("#1D2C3B")
            p_item.space_after = Pt(4)

        # 4. Add Icon image if provided
        if shape_bp.icon_name:
            icon_path = find_icon_file(shape_bp.icon_name)
            if icon_path and icon_path.exists():
                icon_size = Inches(0.45)
                icon_left = left + width - icon_size - Inches(0.2)
                icon_top = top + Inches(0.2)
                slide.shapes.add_picture(str(icon_path), icon_left, icon_top, icon_size, icon_size)

    elif shape_bp.type == "chart":
        # Native PowerPoint Chart
        chart_data = CategoryChartData()
        categories = shape_bp.chart_categories or ["Q1", "Q2", "Q3", "Q4"]
        chart_data.categories = categories
        
        series_list = shape_bp.chart_series or [{"name": "Performance", "values": (25, 40, 65, 90)}]
        for s in series_list:
            chart_data.add_series(s.get("name", "Series 1"), s.get("values", (10, 20, 30, 40)))

        chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
        if shape_bp.chart_type == "bar":
            chart_type = XL_CHART_TYPE.BAR_CLUSTERED
        elif shape_bp.chart_type == "line":
            chart_type = XL_CHART_TYPE.LINE
        elif shape_bp.chart_type in ("pie", "donut"):
            chart_type = XL_CHART_TYPE.PIE

        chart_shape = slide.shapes.add_chart(
            chart_type, left, top, width, height, chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = len(series_list) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False

    elif shape_bp.type == "text_box":
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = shape_bp.text or shape_bp.title
        p.font.size = Pt(shape_bp.font_size_pt or 13)
        p.font.bold = shape_bp.font_bold
        p.font.color.rgb = hex_to_rgb(shape_bp.font_color or "#1D2C3B")
